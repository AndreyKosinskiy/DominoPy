import os
from dotenv import load_dotenv
import win32com.client
import openpyxl


load_dotenv()
REPORT_FORM_NAME = 'ReportForm'
notesServer = os.getenv("NOTES_SERVER")
notesFile = os.getenv("NOTES_DB_NAME")
notesPass = os.getenv("NOTES_PASSWORD")

notesViewName = "ExcelAttachmentViewform"

notesSession = win32com.client.Dispatch('Lotus.NotesSession')
notesSession.Initialize(notesPass)
notesDatabase = notesSession.GetDatabase(notesServer,notesFile)

tempDir = 'C:\\temppy'



def makeDocumentGenerator(folderName):
    folder = notesDatabase.GetView(folderName)
    if not folder:
        raise Exception('Folder "%s" not found' % folderName)
    document = folder.GetFirstDocument()
    while document:
        yield document
        document = folder.GetNextDocument(document)

def file_data(filepath):
    """get flat data table  """
    table = [] 
    book = openpyxl.load_workbook(filepath)
    sheet = book.active
    
    for c1 in sheet[sheet.dimensions]:
        row = ''
        for c2 in c1:
            row =row+str(c2.value)+';'
        table.append(row)
    os.remove(filepath)
    return table

def createReport(ouid,data):
    document = notesDatabase.CreateDocument()
    document.replaceItemValue('Form',REPORT_FORM_NAME)
    document.replaceItemValue('OUID',ouid)
    print(data)
    document.replaceItemValue('Result',data)
    document.Save(True,False)

def main():
    for document in makeDocumentGenerator(notesViewName):
        ouid = document.getItemValue('OUID')[0]
        attachments = notesSession.Evaluate("@AttachmentNames", document)
        for attach in attachments:
            embedObj = document.GetAttachment(attach)
            filepath = tempDir +'\\'+ attach
            embedObj.Extractfile(filepath)
            data = file_data(filepath)
            createReport(ouid,data)

main()